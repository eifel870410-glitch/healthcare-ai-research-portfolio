from pathlib import Path
from datetime import datetime

from docx import Document


def markdown_to_docx(markdown_text, subject_id, hits, report_dir):
    """
    將 Markdown 報告轉成 Word。
    每次輸出都加入時間戳，避免覆蓋舊檔案或被 Word 鎖住。
    """

    report_dir = Path(report_dir)
    report_dir.mkdir(parents=True, exist_ok=True)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = report_dir / f"patient_{subject_id}_education_report_{timestamp}.docx"

    doc = Document()

    doc.add_heading("個人化病人衛教報告", level=0)
    doc.add_paragraph(f"Subject ID：{subject_id}")
    doc.add_paragraph(f"產生時間：{datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
    doc.add_paragraph("")

    lines = markdown_text.splitlines()

    for line in lines:
        line = line.strip()

        if not line:
            continue

        if line.startswith("# "):
            doc.add_heading(line.replace("# ", "").strip(), level=1)

        elif line.startswith("## "):
            doc.add_heading(line.replace("## ", "").strip(), level=2)

        elif line.startswith("### "):
            doc.add_heading(line.replace("### ", "").strip(), level=3)

        elif line.startswith("- "):
            doc.add_paragraph(line.replace("- ", "").strip(), style="List Bullet")

        else:
            doc.add_paragraph(line)

    doc.add_page_break()
    doc.add_heading("RAG 檢索來源", level=1)

    for i, h in enumerate(hits[:3], start=1):
        title = h.get("title", f"衛教單 {i}") if isinstance(h, dict) else getattr(h, "title", f"衛教單 {i}")
        source = h.get("source", "") if isinstance(h, dict) else getattr(h, "source", "")
        score = h.get("score", "") if isinstance(h, dict) else getattr(h, "score", "")

        doc.add_heading(f"來源 {i}", level=2)
        doc.add_paragraph(f"標題：{title}")
        doc.add_paragraph(f"來源：{source}")
        doc.add_paragraph(f"相似度分數：{score}")

    doc.save(out_path)

    return str(out_path)
